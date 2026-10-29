"""Plain-pandas helpers for the 'optimize your own experiment' mode: validate variables,
read an uploaded Excel/CSV, and write the Excel file the user takes to the lab."""
import io
import math
import time

import numpy as np
import pandas as pd

TYPES = ("continuous", "integer")
DATA_SHEET, NEXT_SHEET, SETTINGS_SHEET = "All data", "Next experiments", "Settings"


def clean_variables(df):
    """Turn the editable table into a validated list. Returns (variables, errors)."""
    variables, errors, seen = [], [], set()
    for i, row in df.reset_index(drop=True).iterrows():
        vals = [row.get(c) for c in ("name", "min", "max", "type")]
        if all(pd.isna(v) or str(v).strip() == "" for v in vals):
            continue  # untouched blank row
        label = f"Row {i + 1}"
        name = "" if pd.isna(row.get("name")) else str(row["name"]).strip()
        if not name:
            errors.append(f"{label}: give the variable a name.")
            continue
        if name.lower() in seen:
            errors.append(f"{label}: the name '{name}' is used twice.")
            continue
        seen.add(name.lower())
        typ = row.get("type")
        typ = "continuous" if pd.isna(typ) or str(typ).strip() == "" else str(typ).strip().lower()
        if typ not in TYPES:
            errors.append(f"{name}: type must be continuous or integer.")
            continue
        try:
            lo, hi = float(row["min"]), float(row["max"])
        except (TypeError, ValueError):
            errors.append(f"{name}: min and max must be numbers.")
            continue
        if not (math.isfinite(lo) and math.isfinite(hi)):
            errors.append(f"{name}: min and max must be numbers.")
            continue
        if typ == "integer":
            lo, hi = math.ceil(lo), math.floor(hi)
        if hi <= lo:
            errors.append(f"{name}: max must be larger than min" + (" (by at least 1 for an integer)." if typ == "integer" else "."))
            continue
        variables.append({"name": name, "min": float(lo), "max": float(hi), "type": typ})
    if not variables and not errors:
        errors.append("Add at least one variable.")
    return variables, errors


def parse_data(df, variables, response):
    """Match an uploaded table to the variables. Returns a dict with the clean table and arrays.

    Rows with all variables but a blank response are 'pending' (planned or still running).
    """
    names = [v["name"] for v in variables]
    lower = {}
    for c in df.columns:
        lower.setdefault(str(c).strip().lower(), c)
    missing = [n for n in names + [response] if n.lower() not in lower]
    if missing:
        raise ValueError("Your file has no column named: " + ", ".join(f"'{m}'" for m in missing)
                         + ". Columns found: " + ", ".join(f"'{c}'" for c in df.columns) + ".")
    ren = {lower[n.lower()]: n for n in names + [response]}
    d = df.rename(columns=ren).copy()
    d = d.dropna(how="all").reset_index(drop=True)

    num = {n: pd.to_numeric(d[n], errors="coerce") for n in names}
    resp_raw = d[response]
    resp = pd.to_numeric(resp_raw, errors="coerce")
    blank = resp_raw.isna() | (resp_raw.astype(str).str.strip() == "")
    var_ok = pd.concat(num, axis=1).notna().all(axis=1)

    done = var_ok & resp.notna()
    pending = var_ok & blank
    unreadable = var_ok & resp.isna() & ~blank
    incomplete = ~var_ok

    for n in names:
        d[n] = num[n].where(num[n].notna(), d[n])
    d[response] = resp.where(resp.notna(), resp_raw)
    first = names + [response]
    d = d[first + [c for c in d.columns if c not in first]]

    warnings = []
    if incomplete.sum():
        warnings.append(f"{int(incomplete.sum())} row(s) have a missing or non-numeric variable value and are ignored.")
    if unreadable.sum():
        warnings.append(f"{int(unreadable.sum())} row(s) have a response that is not a number and are ignored.")
    for v in variables:
        col = num[v["name"]][done | pending]
        n_out = int(((col < v["min"]) | (col > v["max"])).sum())
        if n_out:
            warnings.append(f"{n_out} value(s) of '{v['name']}' are outside your range {v['min']:g} to {v['max']:g}. "
                            "They are still used for learning.")
    X = np.column_stack([num[n][done] for n in names]) if done.any() else np.empty((0, len(names)))
    P = np.column_stack([num[n][pending] for n in names]) if pending.any() else np.empty((0, len(names)))
    return {"table": d, "X": X.astype(float), "y": resp[done].to_numpy(float), "pending": P.astype(float),
            "n_done": int(done.sum()), "n_pending": int(pending.sum()), "warnings": warnings}


def read_upload(filename, raw):
    """Read an uploaded .csv or .xlsx. Returns (data_df, settings_or_None)."""
    if filename.lower().endswith(".csv"):
        return pd.read_csv(io.BytesIO(raw)), None
    sheets = pd.read_excel(io.BytesIO(raw), sheet_name=None)
    if not sheets:
        raise ValueError("The Excel file has no sheets.")
    if DATA_SHEET in sheets:
        data = sheets[DATA_SHEET]
    else:
        data = next((df for name, df in sheets.items() if name not in (SETTINGS_SHEET, NEXT_SHEET)),
                    next(iter(sheets.values())))
    return data, parse_settings(sheets.get(SETTINGS_SHEET))


def parse_settings(df):
    if df is None or "item" not in df.columns:
        return None
    try:
        v = df[df["item"] == "variable"]
        variables = pd.DataFrame({"name": v["name"], "min": pd.to_numeric(v["min"]),
                                  "max": pd.to_numeric(v["max"]), "type": v["type"]}).reset_index(drop=True)
        get = lambda k: df.loc[df["item"] == k, "value"].astype(str).iloc[0]
        return {"variables": variables, "response": get("response"), "goal": get("goal").capitalize()}
    except Exception:
        return None


def settings_table(variables, response, maximize):
    rows = [{"item": "variable", "name": v["name"], "min": v["min"], "max": v["max"], "type": v["type"], "value": None}
            for v in variables]
    rows += [{"item": "response", "name": None, "min": None, "max": None, "type": None, "value": response},
             {"item": "goal", "name": None, "min": None, "max": None, "type": None,
              "value": "maximize" if maximize else "minimize"},
             {"item": "created", "name": None, "min": None, "max": None, "type": None,
              "value": time.strftime("%Y-%m-%d %H:%M")}]
    return pd.DataFrame(rows)


def new_rows_table(variables, response, points):
    names = [v["name"] for v in variables]
    t = pd.DataFrame(points, columns=names)
    for v in variables:
        if v["type"] == "integer":
            t[v["name"]] = t[v["name"]].astype(int)
    t[response] = np.nan
    return t


def build_workbook(variables, response, maximize, existing, suggestions):
    """Excel bytes: next experiments (with predictions), all data (existing + blank new rows), settings."""
    from openpyxl.styles import Font, PatternFill
    names = [v["name"] for v in variables]
    new = suggestions[names + [response]]
    base = existing if existing is not None else pd.DataFrame(columns=names + [response])
    all_data = pd.concat([base, new], ignore_index=True) if len(base) else new.copy()
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        suggestions.to_excel(xw, sheet_name=NEXT_SHEET, index=False)
        all_data.to_excel(xw, sheet_name=DATA_SHEET, index=False)
        settings_table(variables, response, maximize).to_excel(xw, sheet_name=SETTINGS_SHEET, index=False)
        for ws in xw.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2D6A4F")
            ws.freeze_panes = "A2"
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 60)
        ws = xw.book[DATA_SHEET]
        fill = PatternFill("solid", fgColor="FFF3CD")
        resp_col = list(all_data.columns).index(response) + 1
        for r in range(len(base) + 2, len(all_data) + 2):
            ws.cell(row=r, column=resp_col).fill = fill  # highlight the cells to fill in
    return buf.getvalue()
